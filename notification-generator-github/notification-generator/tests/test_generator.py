import tempfile
import unittest
from datetime import datetime
from pathlib import Path

import pandas as pd
from docx import Document

from notification_generator import generate, replace_in_paragraph


COLUMNS = ["№", "УПД", "Дата", "Сумма", "Спецификация", "Договор", "Примечание", "Транш", "Оплата"]


class GeneratorTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.excel = self.root / "data.xlsx"
        self.template = self.root / "template.docx"
        self.out = self.root / "output"
        self.rows = [
            [101, "000042", datetime(2025, 12, 10), 1200.5, "СП-1", "Договор DEMO-1", "", datetime(2025, 12, 15), datetime(2026, 1, 15)],
            [101, 43, datetime(2025, 12, 11), 300, "СП-2", "Договор DEMO-1", "", datetime(2025, 12, 15), datetime(2026, 1, 15)],
            [102, 100, datetime(2025, 12, 12), 55.25, "СП-3", "Договор DEMO-2", "", datetime(2025, 12, 16), datetime(2026, 1, 16)],
        ]
        self.save_excel()
        doc = Document()
        doc.add_table(rows=1, cols=1).cell(0, 0).text = "Исх. №{{NOTIFICATION_NUMBER}} от {{NOTIFICATION_DATE}}"
        doc.add_paragraph("Договор: {{CONTRACT_NUMBER}}")
        table = doc.add_table(rows=1, cols=6)
        for i in range(6):
            table.rows[0].cells[i].text = str(i)
        doc.save(self.template)

    def save_excel(self):
        pd.DataFrame(self.rows, columns=COLUMNS).to_excel(self.excel, index=False)

    def test_generates_groups_and_exact_totals(self):
        self.assertEqual(generate(self.excel, self.template, self.out), 2)
        doc = Document(self.out / "Уведомление_№101.docx")
        self.assertIn("№101", doc.tables[0].cell(0, 0).text)
        self.assertIn("Договор DEMO-1", doc.paragraphs[0].text)
        self.assertEqual(len(doc.tables[1].rows), 4)
        self.assertEqual(doc.tables[1].rows[1].cells[1].text, "000042")
        self.assertEqual(doc.tables[1].rows[-1].cells[3].text, "1 500.50")
        self.assertEqual(len(list(self.out.glob("*.docx"))), 2)

    def test_conflicting_numbers_fail_before_writing(self):
        self.rows[1][0] = 103
        self.save_excel()
        with self.assertRaisesRegex(ValueError, "разные номера"):
            generate(self.excel, self.template, self.out)
        self.assertFalse(self.out.exists())

    def test_duplicate_filename_and_existing_output_fail(self):
        self.rows[2][0] = 101
        self.save_excel()
        with self.assertRaisesRegex(ValueError, "Одинаковый номер"):
            generate(self.excel, self.template, self.out)
        self.rows[2][0] = 102
        self.save_excel()
        generate(self.excel, self.template, self.out)
        with self.assertRaisesRegex(ValueError, "уже существует"):
            generate(self.excel, self.template, self.out)

    def test_missing_template_marker_fails(self):
        doc = Document(self.template)
        doc.paragraphs[0].text = "Нет метки договора"
        doc.save(self.template)
        with self.assertRaisesRegex(ValueError, "CONTRACT_NUMBER"):
            generate(self.excel, self.template, self.out)

    def test_marker_split_across_runs(self):
        doc = Document()
        p = doc.add_paragraph("до ")
        p.add_run("{{NOTIFI")
        p.add_run("CATION_NUMBER}} после")
        self.assertEqual(replace_in_paragraph(p, "{{NOTIFICATION_NUMBER}}", "101"), 1)
        self.assertEqual(p.text, "до 101 после")


if __name__ == "__main__":
    unittest.main()
