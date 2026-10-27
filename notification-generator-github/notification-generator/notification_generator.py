#!/usr/bin/env python3
"""Generate Word notifications from an Excel worksheet and a Word template."""

import argparse
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


MONTHS = (
    "", "января", "февраля", "марта", "апреля", "мая", "июня",
    "июля", "августа", "сентября", "октября", "ноября", "декабря",
)
NUMBER_MARKER = "{{NOTIFICATION_NUMBER}}"
DATE_MARKER = "{{NOTIFICATION_DATE}}"
CONTRACT_MARKER = "{{CONTRACT_NUMBER}}"


def required(value, label, row):
    if pd.isna(value) or str(value).strip() == "":
        raise ValueError(f"Строка Excel {row}: пустое поле «{label}»")
    return value


def parse_date(value, label, row):
    required(value, label, row)
    try:
        result = pd.to_datetime(value, dayfirst=True, errors="raise")
        if pd.isna(result):
            raise ValueError("пустая дата")
        return result.date()
    except (ValueError, TypeError, OverflowError) as exc:
        raise ValueError(f"Строка Excel {row}: неверная дата «{label}»: {value}") from exc


def parse_number(value, row):
    required(value, "№ уведомления", row)
    try:
        number = Decimal(str(value))
        if not number.is_finite() or number != number.to_integral_value() or number < 0:
            raise ValueError("номер должен быть целым неотрицательным числом")
        return int(number)
    except (InvalidOperation, ValueError) as exc:
        raise ValueError(f"Строка Excel {row}: неверный номер уведомления: {value}") from exc


def parse_amount(value, row):
    required(value, "Сумма документа", row)
    try:
        amount = Decimal(str(value).replace(" ", "").replace("\u00a0", "").replace(",", "."))
        if not amount.is_finite() or amount < 0 or amount.as_tuple().exponent < -2:
            raise ValueError("ожидается неотрицательная сумма с точностью до копейки")
        return amount
    except (InvalidOperation, ValueError) as exc:
        raise ValueError(f"Строка Excel {row}: неверная сумма: {value}") from exc


def format_amount(amount):
    return f"{amount:,.2f}".replace(",", " ")


def format_upd(value, row):
    required(value, "УПД", row)
    if isinstance(value, (int, float)):
        number = Decimal(str(value))
        if not number.is_finite() or number != number.to_integral_value() or number < 0:
            raise ValueError(f"Строка Excel {row}: неверный номер УПД: {value}")
        return str(int(number)).zfill(6)
    return str(value).strip().zfill(6)


def replace_in_paragraph(paragraph, marker, value):
    """Replace a marker across runs, preserving text outside it and first-run styling."""
    replaced = 0
    while marker in paragraph.text:
        start = paragraph.text.index(marker)
        end = start + len(marker)
        offset = 0
        for run in paragraph.runs:
            original = run.text
            next_offset = offset + len(original)
            if next_offset > start and offset < end:
                prefix = original[:max(0, start - offset)]
                suffix = original[max(0, end - offset):]
                run.text = prefix + (value if offset <= start < next_offset else "") + suffix
            offset = next_offset
        replaced += 1
    return replaced


def paragraphs_in_document(doc):
    yield from doc.paragraphs
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs


def load_notifications(excel_path):
    df = pd.read_excel(excel_path, dtype=object)
    if df.shape[1] < 9:
        raise ValueError("В Excel требуется не менее 9 столбцов (A–I)")
    if df.empty:
        raise ValueError("Excel не содержит строк данных")

    groups = {}
    for index, row in enumerate(df.itertuples(index=False, name=None), start=2):
        if all(pd.isna(item) for item in row[:9]):
            continue
        number = parse_number(row[0], index)
        upd = format_upd(row[1], index)
        doc_date = parse_date(row[2], "Дата документа", index)
        amount = parse_amount(row[3], index)
        spec = "" if pd.isna(row[4]) else str(row[4]).strip()
        contract = str(required(row[5], "№ договора", index)).strip()
        tranche = parse_date(row[7], "Дата выдачи транша", index)
        payment = parse_date(row[8], "Дата оплаты по договору", index)
        key = (contract, tranche)
        group = groups.setdefault(key, {"number": number, "items": [], "total": Decimal("0")})
        if group["number"] != number:
            raise ValueError(f"Строка Excel {index}: у одной группы разные номера уведомления")
        group["items"].append((upd, doc_date.strftime("%d.%m.%Y"), amount, spec, payment.strftime("%d.%m.%Y")))
        group["total"] += amount

    if not groups:
        raise ValueError("Нет строк для создания уведомлений")
    numbers = [group["number"] for group in groups.values()]
    if len(numbers) != len(set(numbers)):
        raise ValueError("Одинаковый номер уведомления назначен разным договорам или датам")
    return sorted(groups.items(), key=lambda item: item[1]["number"])


def validate_template(template_path):
    doc = Document(template_path)
    if len(doc.tables) < 2 or len(doc.tables[1].columns) != 6 or len(doc.tables[1].rows) < 1:
        raise ValueError("Шаблон должен содержать две таблицы; во второй нужны заголовок и 6 столбцов")
    text = "\n".join(p.text for p in paragraphs_in_document(doc))
    for marker in (NUMBER_MARKER, DATE_MARKER, CONTRACT_MARKER):
        if marker not in text:
            raise ValueError(f"В шаблоне не найдена метка {marker}")


def generate_document(template_path, target, key, group):
    contract, tranche = key
    doc = Document(template_path)
    replacements = {
        NUMBER_MARKER: str(group["number"]),
        DATE_MARKER: f"«{tranche.day:02d}» {MONTHS[tranche.month]} {tranche.year} г.",
        CONTRACT_MARKER: contract,
    }
    for paragraph in paragraphs_in_document(doc):
        for marker, value in replacements.items():
            replace_in_paragraph(paragraph, marker, value)

    table = doc.tables[1]
    for row in list(table.rows)[1:]:
        table._tbl.remove(row._tr)
    for i, (upd, doc_date, amount, spec, payment) in enumerate(group["items"], 1):
        add_row(table, (str(i), upd, doc_date, format_amount(amount), spec, payment))
    add_row(table, ("", "", "Итого:", format_amount(group["total"]), "", ""), bold=True)
    doc.save(target)


def add_row(table, values, bold=False):
    row = table.add_row()
    for cell, value in zip(row.cells, values):
        cell.text = str(value)
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                run.font.name = "Arial"
                run.font.size = Pt(8)
                run.bold = bold


def generate(excel_path, template_path, output_dir):
    excel_path, template_path, output_dir = map(Path, (excel_path, template_path, output_dir))
    if not excel_path.is_file() or not template_path.is_file():
        raise ValueError("Проверьте пути к Excel и Word-шаблону")
    groups = load_notifications(excel_path)
    validate_template(template_path)
    for key, group in groups:
        target = output_dir / f"Уведомление_№{group['number']}.docx"
        if target.exists():
            raise ValueError(f"Выходной файл уже существует: {target}. Выберите пустую папку")
    if (output_dir / "report.txt").exists():
        raise ValueError("Отчёт уже существует в выходной папке. Выберите пустую папку")
    output_dir.mkdir(parents=True, exist_ok=True)
    for key, group in groups:
        generate_document(template_path, output_dir / f"Уведомление_№{group['number']}.docx", key, group)
    lines = [f"Дата: {datetime.now():%d.%m.%Y %H:%M:%S}", f"Создано уведомлений: {len(groups)}", ""]
    for (contract, tranche), group in groups:
        lines.append(f"№{group['number']} от {tranche:%d.%m.%Y}; договор: {contract}; документов: {len(group['items'])}; сумма: {format_amount(group['total'])}")
    (output_dir / "report.txt").write_text("\n".join(lines) + "\n", encoding="utf-8")
    return len(groups)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("excel_file", help="Исходная таблица .xlsx")
    parser.add_argument("template_file", help="Шаблон .docx с метками {{...}}")
    parser.add_argument("output_dir", nargs="?", default="notifications")
    args = parser.parse_args()
    try:
        count = generate(args.excel_file, args.template_file, args.output_dir)
    except (ValueError, OSError, ImportError) as exc:
        parser.exit(1, f"Ошибка: {exc}\n")
    print(f"Создано уведомлений: {count}. Результат: {args.output_dir}")


if __name__ == "__main__":
    main()
