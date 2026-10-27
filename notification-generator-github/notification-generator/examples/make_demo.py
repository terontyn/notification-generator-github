"""Create fictional input files for a local smoke test."""

from datetime import datetime
from pathlib import Path

import pandas as pd
from docx import Document


root = Path(__file__).resolve().parent / "generated"
root.mkdir(parents=True, exist_ok=True)

rows = [
    [101, "000042", datetime(2025, 12, 10), 1200.50, "СП-1", "Договор DEMO-1", "", datetime(2025, 12, 15), datetime(2026, 1, 15)],
    [101, 43, datetime(2025, 12, 11), 300, "СП-2", "Договор DEMO-1", "", datetime(2025, 12, 15), datetime(2026, 1, 15)],
    [102, 100, datetime(2025, 12, 12), 55.25, "СП-3", "Договор DEMO-2", "", datetime(2025, 12, 16), datetime(2026, 1, 16)],
]
columns = ["№ уведомления", "УПД", "Дата документа", "Сумма документа", "№ спецификации", "№ договора", "Примечание", "Дата транша", "Дата оплаты"]
pd.DataFrame(rows, columns=columns).to_excel(root / "data.xlsx", index=False)

doc = Document()
header = doc.add_table(rows=1, cols=1)
header.cell(0, 0).text = "Исх. №{{NOTIFICATION_NUMBER}} от {{NOTIFICATION_DATE}}"
doc.add_paragraph("Уведомление по договору {{CONTRACT_NUMBER}}.")
table = doc.add_table(rows=1, cols=6)
for cell, title in zip(table.rows[0].cells, ("№", "УПД", "Дата документа", "Сумма", "Спецификация", "Дата оплаты")):
    cell.text = title
doc.save(root / "template.docx")
print(f"Созданы вымышленные файлы в {root}")
